from pathlib import Path

from docx import Document


root = Path(r"E:\Prashanth\finance-tracker v2")
source = root / "PROJECT_REPORT_RupeeFlow_Current.docx"
output = root / "PROJECT_REPORT_RupeeFlow_Final.docx"
doc = Document(source)


def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell(cell, text):
    set_text(cell.paragraphs[0], text)


set_text(doc.paragraphs[365], "Final verification covered code quality, production compilation, module calculations, statement parsing, duplicate identity, recurring generation, route availability, and OTP service health. No retained feature failed these automated checks.")
set_text(doc.paragraphs[367], "Automated module checks covered totals, budget states, savings progress, trends, forecasts, duplicate fingerprints, recurring schedules, and INR formatting; all assertions passed.")
set_text(doc.paragraphs[368], "Statement parsing was validated with two real supported bank-statement samples: 127 transactions from the May statement and 235 transactions from the earlier statement were extracted successfully.")
set_text(doc.paragraphs[369], "OTP service health confirmed valid SMTP configuration; authenticated Gmail SMTP accepted a controlled delivery test, while invalid email, OTP, and reset-token requests were rejected correctly.")
set_text(doc.paragraphs[370], "Application routes for dashboard, transactions, budgets, goals, recurring rules, reports, and profile returned successfully through the development server.")
set_text(doc.paragraphs[371], "ESLint completed without errors and Vite produced a successful optimized build from 611 transformed modules. The remaining large-chunk message is a performance advisory, not a functional failure.")

tests = doc.tables[16].rows
set_cell(tests[8].cells[2], "Supported samples extract 127 and 235 transactions; records are categorized and saved")
set_cell(tests[9].cells[2], "Existing transactions are skipped by duplicate fingerprints")
set_cell(tests[17].cells[2], "All fields and action buttons remain scrollable and reachable")

doc.core_properties.title = "RupeeFlow Finance Tracker - Final Project Report"
doc.core_properties.subject = "Final report synchronized with current application and verification results"
doc.save(output)
print(output)
