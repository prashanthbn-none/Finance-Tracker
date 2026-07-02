from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt


ROOT = Path(r"E:\Prashanth\finance-tracker v2")
BASE = ROOT / "PROJECT_REPORT_RupeeFlow_Academic_Final.docx"
SHOT_ROOT = ROOT / "report_screenshots"

REPORTS = [
    {
        "kind": "individual", "folder": "prashanth", "name": "Prashanth B N", "id": "U18ID23S0050",
        "output": ROOT / "RupeeFlow_Project_Report_Prashanth_BN.docx",
    },
    {
        "kind": "individual", "folder": "lochan", "name": "Lochan M", "id": "U18ID23S0035",
        "output": ROOT / "RupeeFlow_Project_Report_Lochan_M.docx",
    },
    {
        "kind": "individual", "folder": "koushik", "name": "Koushik R", "id": "U18ID23S0051",
        "output": ROOT / "RupeeFlow_Project_Report_Koushik_R.docx",
    },
    {
        "kind": "group", "folder": "prashanth",
        "members": [
            ("Prashanth B N", "U18ID23S0050"),
            ("Lochan M", "U18ID23S0035"),
            ("Koushik R", "U18ID23S0051"),
        ],
        "output": ROOT / "RupeeFlow_Project_Report_Group.docx",
    },
]

SHOTS = [
    ("dashboard.png", "Dashboard"),
    ("transactions.png", "Transaction Management"),
    ("budgets.png", "Budget Management"),
    ("goals.png", "Savings Goals"),
    ("reports.png", "Reports and Analytics"),
    ("profile.png", "Profile and INR Settings"),
]

APPENDICES = [
    ("A.1 Verified Project Structure", [
        "RupeeFlow uses a compact client-plus-service structure. The finance-tracker/src directory contains route-level pages, reusable components, Context stores, hooks, styles, and pure library modules. The server directory contains the narrowly scoped OTP email service. Vite configuration connects frontend API requests to that service during development.",
        "Important client files include App.jsx for route composition, AppShell.jsx for navigation, DataContext.jsx for collection operations, AuthContext.jsx for local account access, storage.js for IndexedDB persistence, selectors.js for derived finance metrics, recurring.js for schedule materialisation, statementImport.js for text parsing, and export.js for CSV, Excel, and print workflows.",
        "This structure follows separation of concerns: pages manage interaction, Context providers coordinate state, library files implement reusable logic, and the storage adapter owns persistence. The report describes this actual structure rather than a hypothetical Express or multi-database architecture."
    ]),
    ("A.2 Application Startup and Routing", [
        "The application starts in main.jsx, mounts React in strict mode, and renders App. App composes BrowserRouter, SettingsProvider, AuthProvider, and a Gate component. Gate waits for account restoration, displays the authentication screen when no local session exists, and otherwise mounts DataProvider and protected finance routes.",
        "Routes cover dashboard, income, expenses, transactions, budgets, savings goals, recurring transactions, reports, and profile settings. AppShell supplies persistent navigation, theme control, notifications, and the account menu. Unknown paths redirect to the dashboard.",
        "This design keeps route definitions centralized and allows every finance screen to share authentication, settings, and data state without prop drilling. Route availability was checked through the development server during final verification."
    ]),
    ("A.3 Authentication and Local Account Flow", [
        "Registration accepts name, email, and password. AuthContext checks duplicate emails, computes a SHA-256 digest through the Web Crypto API, stores the local account record, and establishes a browser session. Login recomputes the digest and compares it with the stored value before restoring the public user object.",
        "This mechanism is suitable only for the academic local-first demonstration. It does not use a production identity database, salted Argon2 or bcrypt hashes, secure HTTP-only cookies, or server-enforced authorization. The interface and report explicitly retain this limitation.",
        "A future hosted version should move account creation, password hashing, session issuance, authorization, and audit controls to a secure backend while preserving the current frontend Context interface."
    ]),
    ("A.4 Email OTP Password Recovery", [
        "Password recovery uses three operations: request a six-digit code, verify that code, and consume a one-time reset token. The browser first confirms that the email belongs to a local account, then calls the Node.js service through /api/auth endpoints.",
        "The service stores salted OTP hashes in memory, applies a ten-minute expiry, permits five verification attempts, and enforces a sixty-second resend delay. Successful verification replaces the OTP challenge with a short-lived one-time reset token. Nodemailer sends the code through authenticated SMTP using credentials loaded from a private environment file.",
        "No SMTP password is placed in frontend code or committed to source control. Delivery can still be affected by provider spam filtering, service availability, and local configuration."
    ]),
    ("A.5 IndexedDB Storage Adapter", [
        "Financial records are stored in the browser database named rupeeflow. A collections object store keeps records keyed by user and collection name. Transactions, budgets, goals, and recurring rules are therefore separated for each local account without requiring multiple physical object stores.",
        "DataContext uses createStorage(userId) and calls list, write, getMeta, setMeta, and clearAll. IndexedDbStorageAdapter is preferred; LocalStorageAdapter remains a compatibility fallback and migration source. This boundary allows a future hosted storage implementation without redesigning feature pages.",
        "Metadata records hold dismissed alerts, migration status, and custom categories. Clearing a user's finance data deletes only keys within that user's namespace."
    ]),
    ("A.6 Transaction Record Lifecycle", [
        "A transaction contains id, type, amount, category, date, note, and an optional recurringId. TransactionModal validates positive amount, category, and date before calling DataContext. The Context creates an identifier and persists the complete updated collection.",
        "Editing builds a candidate record, performs duplicate comparison excluding the current identifier, and then persists the patch. Deletion removes the selected identifier. Bulk import validates required values and assigns identifiers only to accepted records.",
        "All display formatting uses the INR formatter with the en-IN locale. Income and expense screens share one editor while retaining their relevant category lists."
    ]),
    ("A.7 Duplicate Prevention Algorithm", [
        "Duplicate prevention uses a deterministic fingerprint composed of date, transaction type, amount rounded to two decimal places, category, and normalized note text. Normalization converts note text to lowercase, collapses repeated whitespace, and trims outer spaces.",
        "The same fingerprint function is used for manual creation, editing, recurring generation, CSV/PDF import, and duplicate checks within a single imported batch. A Set provides efficient membership testing when many statement rows are processed.",
        "This approach prevents accidental repeated imports while allowing genuinely different transactions. Two identical real purchases on the same date with the same note and amount may require manual distinction, which is a known trade-off of deterministic duplicate rules."
    ]),
    ("A.8 PDF and CSV Statement Processing", [
        "The transaction screen accepts CSV and unlocked selectable-text PDF files. PDF.js extracts text inside the browser, preserving the privacy-first architecture. statementImport.js then evaluates supported row patterns, parses dates and debit or credit amounts, normalizes narration, and converts valid rows into transaction candidates.",
        "Rule-based keyword matching assigns India-relevant categories such as food, transportation, utilities, EMI, rent, healthcare, education, shopping, salary, and other income. Rows that cannot be validated are skipped instead of being stored with uncertain values.",
        "Testing extracted 127 transactions from one supplied statement and 235 from another supported statement. The system does not claim universal support for every Indian bank layout, scanned image PDF, or password-protected statement."
    ]),
    ("A.9 Budget Evaluation", [
        "A budget record combines month, expense category, and limit. budgetStatus derives matching monthly spend from transactions, computes utilization ratio, remaining value, and state. Progress bars use normal, near-limit, and over-limit colors so the user can interpret status quickly.",
        "One budget is maintained for a category in a given month. Saving an existing month-category pair updates its limit rather than creating a second budget. Notifications are derived from current records and can report near or exceeded thresholds.",
        "Custom expense categories can also receive budgets. This keeps budgeting practical for household-specific items not covered by predefined categories."
    ]),
    ("A.10 Savings Goals", [
        "Savings goals store a name, target amount, saved amount, and optional target date. Users can create goals, edit goal details, contribute funds, withdraw funds, and delete goals. Contribution logic prevents saved value from falling below zero.",
        "The interface derives completion percentage, remaining amount, completion state, and days remaining. Circular progress graphics communicate status without replacing exact INR values.",
        "Goals are planning records rather than bank accounts. Contributions change the application's tracking value only and do not transfer money through any payment or banking service."
    ]),
    ("A.11 Recurring Transaction Scheduler", [
        "Recurring rules describe type, amount, category, frequency, start date, note, last run, and paused state. Weekly, monthly, and yearly frequencies are supported. materializeDue calculates periods that should exist up to the current date and returns generated transactions plus updated rules.",
        "Materialization occurs when data loads and when a new rule is created. Generated records include recurringId and pass through duplicate fingerprint checks. Rules may be paused, resumed, edited, or deleted without removing transactions already generated.",
        "The scheduler is browser-driven, so records are materialized when the application next opens rather than by a continuously running cloud job."
    ]),
    ("A.12 Dashboard and Derived Analytics", [
        "Dashboard metrics are derived from shared transaction, budget, goal, and recurring collections. Selectors compute balance, current-month income and spending, savings rate, top category, budget status, recent activity, bills due, and a thirty-day projection.",
        "Recharts renders income-versus-expense trends and category distribution. Pure selector functions keep financial calculations separate from visual components and permit direct module testing with controlled inputs.",
        "Every dashboard value updates after manual entry, edit, deletion, import, recurring generation, backup restoration, or budget change because pages consume the same DataContext state."
    ]),
    ("A.13 Reports and Export", [
        "Reports provide three-, six-, and twelve-month views with total income, expenses, net result, average monthly spend, monthly comparison, savings trend, and category breakdown. CSV and Excel-compatible exports create files from structured rows.",
        "PDF output uses the browser print system. A hidden same-page iframe receives a cloned report node and styles. The helper waits for document fonts and rendered content before opening the print dialog, preventing the previously observed blank preview.",
        "Print-specific rules remove interactive controls, suppress shadows, keep cards together where possible, and apply white background and readable text for paper or PDF output."
    ]),
    ("A.14 Responsive Modal and Form Design", [
        "Modal windows render through a React portal directly under document.body. This prevents transformed or animated route containers from clipping fixed overlays. The modal uses viewport-aware maximum height, fixed header, internal scrolling, responsive width, and wrapping actions.",
        "Forms remain usable on smaller screens because inputs retain full width, action rows wrap, and long content scrolls inside the dialog. Escape closes the modal, outside clicks close only when the overlay itself is selected, and body scrolling is restored after closure.",
        "Transaction, budget, goal, and recurring forms all reuse these layout rules, fixing trimmed fields and inaccessible submit buttons."
    ]),
    ("A.15 Custom Category Management", [
        "CategorySelect combines predefined lists with per-user customCategories metadata. A New button requests a category name, normalizes repeated spaces, rejects empty values, and performs case-insensitive duplicate comparison.",
        "Custom income and expense categories become immediately selectable in transaction and recurring forms. Expense categories can also be used for budgets. Because categories are saved per user, accounts sharing a browser do not receive each other's custom lists.",
        "Unknown category colors use a neutral fallback, ensuring charts and transaction markers remain visible without requiring source-code changes for each user-defined category."
    ]),
    ("A.16 Security Review and Privacy Boundaries", [
        "Financial collections remain local to the browser and are not sent to the OTP service. React escapes ordinary text rendering, input validation checks expected fields and ranges, and statement parsing rejects unsupported rows. Environment files containing SMTP credentials are excluded from version control.",
        "Local storage does not equal complete security. Anyone with access to the same unlocked operating-system account and browser profile may access browser data. IndexedDB records are not application-level encrypted, and local authentication is not suitable for public production deployment.",
        "A production version requires HTTPS, server-side authentication and authorization, secure password hashing, protected sessions, rate limiting, audit logging, monitoring, encrypted backups, and formal dependency maintenance."
    ]),
    ("A.17 Verification Matrix and Evidence", [
        "Final code verification included ESLint, optimized Vite compilation, route availability, OTP health, selector calculations, INR formatting, recurring materialization, duplicate fingerprints, and statement parsing. The optimized build transformed 611 modules successfully; the large-chunk notice remained an advisory rather than a functional failure.",
        "Manual visual checks covered login, dashboard, transactions, budgets, goals, recurring rules, reports, profile settings, light theme, dark theme, modal scrolling, and print preparation. Report screenshots use representative demonstration data and identify the relevant student or logged-in demonstration user.",
        "Evidence must be interpreted within project scope. Successful tests confirm tested paths and supplied statement samples, not every bank format, every browser version, or production-scale security."
    ]),
    ("A.18 Deployment, Maintenance, and Extension", [
        "Development uses npm run dev to start Vite and the OTP service together. Vite proxies /api requests to the local Node service. Production compilation uses npm run build, while npm run lint performs static quality checks. Root convenience scripts delegate to the real finance-tracker application folder.",
        "Repository maintenance should keep .env ignored, commit .env.example without secrets, review dependency updates, run lint and build before pushing, and avoid committing generated reports or personal bank statements unless explicitly required and sanitized.",
        "Future extensions include bank-specific parsers for major Indian banks, OCR for scanned PDFs, confidence-based import review, secure cloud synchronization, investment and EMI modules, shared household accounts, and consent-based regulated bank integrations."
    ]),
]


def set_text(paragraph, text):
    if paragraph.runs:
        first = paragraph.runs[0]
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
        first.text = text
    else:
        paragraph.add_run(text)


def replace_all(doc, old, new):
    for paragraph in doc.paragraphs:
        if old in paragraph.text:
            set_text(paragraph, paragraph.text.replace(old, new))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        set_text(paragraph, paragraph.text.replace(old, new))


def find_para(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(text)


def normalize_readability(doc):
    """Use readable academic typography and remove wasteful body page breaks."""
    for style_name, size in (
        ("Normal", 11),
        ("No Spacing", 11),
        ("List Paragraph", 11),
        ("Heading 3", 11.5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)

    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.15
    normal.space_after = Pt(4)
    doc.styles["List Paragraph"].paragraph_format.line_spacing = 1.08

    # Front matter deliberately uses page breaks. Body chapters should flow
    # naturally, avoiding half-empty pages caused by forced chapter starts.
    for paragraph in doc.paragraphs[70:]:
        for br in paragraph._element.xpath('.//w:br[@w:type="page"]'):
            br.getparent().remove(br)

    # Tables need compact but readable type; preserve code samples separately.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)


def replace_screenshot_section(doc, folder, identity):
    start = find_para(doc, "13. SCREENSHOTS")
    end = find_para(doc, "14. RESULTS, APPLICATIONS & CONCLUSION")
    node = start._element.getnext()
    while node is not None and node is not end._element:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt

    intro = end.insert_paragraph_before(
        "The following screenshots show the completed RupeeFlow application using representative demonstration data."
    )
    intro.style = doc.styles["Normal"]

    for index, (filename, label) in enumerate(SHOTS, 1):
        caption = end.insert_paragraph_before()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(3)
        run = caption.add_run(f"Figure 13.{index}: {identity} - {label}")
        run.bold = True
        run.font.size = Pt(10)

        image_p = end.insert_paragraph_before()
        image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_p.paragraph_format.space_after = Pt(5)
        image_p.add_run().add_picture(str(SHOT_ROOT / folder / filename), width=Inches(5.25))
        if index in (2, 4):
            image_p.add_run().add_break(WD_BREAK.PAGE)


def add_coding_chapter(doc):
    testing = find_para(doc, "12. TESTING & VALIDATION")
    title = testing.insert_paragraph_before()
    title.style = doc.styles["Heading 1"]
    title.add_run("11. CODING AND IMPLEMENTATION")
    title.paragraph_format.page_break_before = True

    for number, (heading, paragraphs) in enumerate(APPENDICES, 1):
        section = testing.insert_paragraph_before()
        section.style = doc.styles["Heading 2"]
        section.add_run(f"11.{number} {heading.split(' ', 1)[1]}")
        section.paragraph_format.page_break_before = False
        for text in paragraphs:
            paragraph = testing.insert_paragraph_before(text)
            paragraph.style = doc.styles["Normal"]
            paragraph.alignment = 3

    code_files = [
        ("Application Routes", ROOT / "finance-tracker/src/App.jsx"),
        ("Authentication Context", ROOT / "finance-tracker/src/store/AuthContext.jsx"),
        ("Finance Data Context", ROOT / "finance-tracker/src/store/DataContext.jsx"),
        ("IndexedDB Adapter", ROOT / "finance-tracker/src/lib/storage.js"),
        ("Duplicate Fingerprint", ROOT / "finance-tracker/src/lib/domain.js"),
        ("Statement Import", ROOT / "finance-tracker/src/lib/statementImport.js"),
        ("Recurring Scheduler", ROOT / "finance-tracker/src/lib/recurring.js"),
        ("Derived Selectors", ROOT / "finance-tracker/src/lib/selectors.js"),
        ("Export and Printing", ROOT / "finance-tracker/src/lib/export.js"),
        ("Shared Transaction Form", ROOT / "finance-tracker/src/components/shared.jsx"),
        ("Responsive Modal", ROOT / "finance-tracker/src/components/ui.jsx"),
        ("OTP Email Service", ROOT / "finance-tracker/server/otpServer.js"),
    ]
    for offset, (label, path) in enumerate(code_files, 1):
        heading = testing.insert_paragraph_before()
        heading.style = doc.styles["Heading 2"]
        heading.add_run(f"11.{len(APPENDICES) + offset} Code Walkthrough: {label}")
        explanation = testing.insert_paragraph_before(
            f"The following verified excerpt comes from {path.name}. It demonstrates the implemented {label.lower()} logic used by the running RupeeFlow application. The excerpt is included for academic explanation; complete source remains in the project repository."
        )
        explanation.style = doc.styles["Normal"]
        explanation.alignment = 3
        lines = [line.rstrip() for line in path.read_text(encoding="utf-8").splitlines() if line.strip()][:34]
        block = testing.insert_paragraph_before()
        block.paragraph_format.space_before = Pt(4)
        block.paragraph_format.space_after = Pt(8)
        block.paragraph_format.keep_together = True
        block.paragraph_format.left_indent = Inches(0.2)
        block.paragraph_format.right_indent = Inches(0.2)
        block.paragraph_format.line_spacing = 1.0
        run = block.add_run("\n".join(lines))
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)


def prepare_academic_structure(doc, group=False):
    # Separate and align front-matter pages.
    for heading in ("DECLARATION", "ACKNOWLEDGEMENT", "ABSTRACT", "INDEX"):
        paragraph = find_para(doc, heading)
        paragraph.paragraph_format.page_break_before = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    abstract = find_para(doc, "ABSTRACT")
    index_heading = find_para(doc, "INDEX")
    node = abstract._element.getnext()
    while node is not None and node is not index_heading._element:
        if node.tag.endswith('}p'):
            from docx.text.paragraph import Paragraph
            paragraph = Paragraph(node, abstract._parent)
            if paragraph.text.strip():
                paragraph.alignment = 3
                paragraph.paragraph_format.line_spacing = 1.15
        node = node.getnext()
    extra = [
        "RupeeFlow is designed for practical Indian personal-finance scenarios. It supports INR-only values, UPI and EMI-aware categories, custom categories, monthly budgets, savings targets, recurring salary and bills, notifications, multi-period reports, and backup or restore. Its statement workflow processes supported selectable-text PDF and CSV files locally, applies rule-based categorisation, and blocks repeated transaction fingerprints.",
        "The project demonstrates component-based React engineering, centralized Context state, asynchronous IndexedDB persistence, pure financial selector functions, responsive modal design, browser-side PDF processing, and a small Node.js OTP delivery service. The architecture separates presentation, shared state, business rules, persistence, and email verification so each part can be tested and extended independently.",
        "Verification covered production compilation, lint quality, financial calculations, recurring schedules, duplicate prevention, route availability, OTP health, responsive form behaviour, print output, and two supplied selectable-text bank statements. The resulting system provides a practical academic demonstration of privacy-focused finance management while clearly documenting current limits such as local-only storage and bank-format-specific PDF support."
    ]
    for text in reversed(extra):
        p = index_heading.insert_paragraph_before(text)
        p.style = doc.styles["Normal"]
        p.alignment = 3
        p.paragraph_format.line_spacing = 1.15

    # Renumber final chapters to match reference academic structure.
    replacements = {
        "11. TESTING & VALIDATION": "12. TESTING & VALIDATION",
        "12. SCREENSHOTS": "13. SCREENSHOTS",
        "13. ADVANTAGES, APPLICATIONS & CONCLUSION": "14. RESULTS, APPLICATIONS & CONCLUSION",
        "14. LIMITATIONS & FUTURE ENHANCEMENTS": "15. LIMITATIONS & FUTURE ENHANCEMENTS",
        "15. BIBLIOGRAPHY": "16. BIBLIOGRAPHY",
    }
    for old, new in replacements.items():
        set_text(find_para(doc, old), new)
    for old, new in {
        "13.1 Advantages of the System": "14.1 Advantages of the System",
        "13.2 Applications of the System": "14.2 Applications of the System",
        "13.3 Conclusion": "14.3 Conclusion",
        "14.1 Limitations": "15.1 Limitations",
        "14.2 Future Enhancements": "15.2 Future Enhancements",
    }.items():
        set_text(find_para(doc, old), new)

    bibliography = find_para(doc, "16. BIBLIOGRAPHY")
    sequence = 1
    node = bibliography._element.getnext()
    while node is not None:
        if node.tag.endswith('}p'):
            from docx.text.paragraph import Paragraph
            paragraph = Paragraph(node, bibliography._parent)
            if paragraph.text.strip():
                ppr = paragraph._p.get_or_add_pPr()
                numpr = ppr.find(qn("w:numPr"))
                if numpr is not None:
                    ppr.remove(numpr)
                set_text(paragraph, f"{sequence}. {paragraph.text.strip()}")
                sequence += 1
        node = node.getnext()

    index = doc.tables[0]
    labels = {
        "Testing & Validation": ("12.", "Testing & Validation"),
        "Screenshots": ("13.", "Screenshots"),
        "Advantages, Applications & Conclusion": ("14.", "Results, Applications & Conclusion"),
        "Limitations & Future Enhancements": ("15.", "Limitations & Future Enhancements"),
        "Bibliography": ("16.", "Bibliography"),
    }
    for row in index.rows:
        if len(row.cells) > 1 and row.cells[1].text.strip() in labels:
            no, label = labels[row.cells[1].text.strip()]
            row.cells[0].text = no
            row.cells[1].text = label
    testing_row = next(row for row in index.rows if len(row.cells) > 1 and row.cells[1].text.strip() == "Testing & Validation")
    coding_row = index.add_row()
    coding_row.cells[0].text = "11."
    coding_row.cells[1].text = "Coding and Implementation"
    coding_row.cells[2].text = ""
    testing_row._tr.addprevious(coding_row._tr)

    # Visible academic page border, matching supplied reference style.
    for section in doc.sections:
        sect_pr = section._sectPr
        old = sect_pr.find(qn("w:pgBorders"))
        if old is not None:
            sect_pr.remove(old)
        borders = OxmlElement("w:pgBorders")
        borders.set(qn("w:offsetFrom"), "page")
        for edge in ("top", "left", "bottom", "right"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "12")
            border.set(qn("w:space"), "24")
            border.set(qn("w:color"), "000000")
            borders.append(border)
        sect_pr.append(borders)


def set_cell_borders(cell, color="000000", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def build_detailed_index(doc):
    table = doc.tables[0]
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)

    headings = []
    started = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "1. INTRODUCTION":
            started = True
        if not started:
            continue
        if paragraph.style.name in ("Heading 1", "Heading 2") and text and text != "INDEX":
            headings.append((paragraph.style.name, text))

    serial = 1
    for style_name, text in headings:
        row = table.add_row()
        if style_name == "Heading 1":
            row.cells[0].text = f"{serial}."
            row.cells[1].text = text.split(" ", 1)[1] if " " in text else text
            for run in row.cells[1].paragraphs[0].runs:
                run.bold = True
            serial += 1
        else:
            row.cells[0].text = ""
            row.cells[1].text = f"    {text}"
        row.cells[2].text = ""

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_borders(cell, size="8" if row_index == 0 else "5")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.5)

    # Add clear borders to every other table as well.
    for other in doc.tables[1:]:
        for row in other.rows:
            for cell in row.cells:
                set_cell_borders(cell)


for spec in REPORTS:
    doc = Document(BASE)
    normalize_readability(doc)
    if spec["kind"] == "individual":
        replace_all(doc, "Prashanth B N", spec["name"])
        replace_all(doc, "PRASHANTH B N", spec["name"].upper())
        replace_all(doc, "U18ID23S0050", spec["id"])
        replace_all(doc, "U18ID23S0050", spec["id"])
        identity = f'{spec["name"]} ({spec["id"]})'
        doc.core_properties.title = f'RupeeFlow Finance Tracker - {spec["name"]} Project Report'
    else:
        members = spec["members"]
        names = "\n".join(name for name, _ in members)
        ids = "\n".join(identifier for _, identifier in members)
        set_text(doc.paragraphs[14], names)
        set_text(doc.paragraphs[15], ids)
        doc.paragraphs[14].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[15].alignment = WD_ALIGN_PARAGRAPH.CENTER
        # College identity already appears at top of cover. Removing repeated
        # footer lines keeps all three students on page 1 without shrinking type.
        for index in (17, 16):
            element = doc.paragraphs[index]._element
            element.getparent().remove(element)
        find_para(doc, "CERTIFICATE").paragraph_format.page_break_before = False
        declaration_date = next(p for p in doc.paragraphs if p.text.strip().startswith("Date:") and "Prashanth B N" in p.text)
        declaration_place = next(p for p in doc.paragraphs if p.text.strip().startswith("Place: Bangalore") and "U18ID23S0050" in p.text)
        set_text(declaration_date, "Date:")
        set_text(declaration_place, "Place: Bangalore\n\nTeam Members:\nPrashanth B N - U18ID23S0050\nLochan M - U18ID23S0035\nKoushik R - U18ID23S0051")
        declaration_date.alignment = WD_ALIGN_PARAGRAPH.LEFT
        declaration_place.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Heading 1 already starts Abstract on a new page; remove duplicate
        # manual break immediately before it to avoid an empty page.
        abstract = find_para(doc, "ABSTRACT")
        previous = abstract._element.getprevious()
        if previous is not None:
            for br in previous.xpath('.//w:br[@w:type="page"]'):
                br.getparent().remove(br)

        roster_inline = ", ".join(f"{name} [{identifier}]" for name, identifier in members)
        for paragraph in doc.paragraphs:
            if "Certified that the project work entitled" in paragraph.text:
                set_text(paragraph, f"Certified that the project work entitled “RupeeFlow Finance Tracker” is bonafide work carried out by {roster_inline} in partial fulfilment for the award of the Degree of Bachelor of Computer Applications of Bengaluru City University during the year 2025-26. The project report has been approved as it satisfies the academic requirements prescribed for the project work.")
            elif paragraph.text.startswith("I hereby declare"):
                set_text(paragraph, paragraph.text.replace("I hereby declare", "We hereby declare").replace("my own work", "our own work"))
            elif paragraph.text.startswith("is an authentic record of my own work"):
                set_text(paragraph, paragraph.text.replace("my own work", "our own work"))
            elif paragraph.text.strip().startswith("Date: Prashanth B N"):
                set_text(paragraph, "Date:\n" + "\n".join(name for name, _ in members))
            elif paragraph.text.strip().startswith("Place: Bangalore U18ID23S0050"):
                set_text(paragraph, "Place: Bangalore\n" + "\n".join(identifier for _, identifier in members))
            elif paragraph.text.strip() == "Prashanth B N":
                element = paragraph._element
                element.getparent().remove(element)
        replace_all(doc, "my effort", "our effort")
        replace_all(doc, "I express", "We express")
        replace_all(doc, "my sincere", "our sincere")
        replace_all(doc, "providing me", "providing us")
        replace_all(doc, "I am deeply thankful", "We are deeply thankful")
        replace_all(doc, "I extend", "We extend")
        replace_all(doc, "my heartfelt", "our heartfelt")
        replace_all(doc, "encouraged me", "encouraged us")
        replace_all(doc, "I would like", "We would like")
        replace_all(doc, "my friends", "our friends")
        replace_all(doc, "my respected", "our respected")
        replace_all(doc, "my guide", "our guide")
        replace_all(doc, "my parents", "our parents")
        replace_all(doc, "helped me", "helped us")
        replace_all(doc, "throughout the completion of this project", "throughout completion of this group project")
        identity = "Group Project Demonstration - Logged in as Prashanth B N"
        doc.core_properties.title = "RupeeFlow Finance Tracker - Group Project Report"

    prepare_academic_structure(doc, group=spec["kind"] == "group")
    add_coding_chapter(doc)
    replace_screenshot_section(doc, spec["folder"], identity)
    build_detailed_index(doc)
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        element = doc.paragraphs[-1]._element
        element.getparent().remove(element)
    try:
        doc.save(spec["output"])
        print(spec["output"])
    except PermissionError:
        fallback = spec["output"].with_name(spec["output"].stem + "_Fixed.docx")
        doc.save(fallback)
        print(fallback)
